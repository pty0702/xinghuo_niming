
import asyncio
import json
import os
import random
import string
from datetime import datetime

import websockets
from openpyxl import Workbook, load_workbook


WS_URL = "wss://xinghuo1v1.com/ws/"
EXCEL_PATH = "星火匹配记录.xlsx"

# 是否匹配成功记录后主动离开当前匹配
# 建议 True，避免程序结束后占着会话状态
CLOSE_AFTER_LOG = True

# 匹配等待超时时间，单位秒
MATCH_TIMEOUT = 120


def gen_client_id(length=32):
    """
    生成客户端匿名 ID。
    实测网站 sign.id 类似 32 位小写字母 + 数字。
    """
    chars = string.ascii_lowercase + string.digits
    return "".join(random.choice(chars) for _ in range(length))


def ensure_excel(path: str):
    """
    如果 Excel 不存在，则创建并写入表头。
    """
    headers = [
        "记录时间",
        "自身客户端ID(sign.id)",
        "自身昵称",
        "自身性别",
        "对方ID(sel_userid)",
        "对方昵称(sel_userNikename)",
        "对方性别(sel_userSex)",
        "对方年龄(sel_userAge)",
        "对方地址(sel_userAddress)",
        "对方标签(sel_userTag)",
        "原始code15 JSON",
    ]

    if not os.path.exists(path):
        wb = Workbook()
        ws = wb.active
        ws.title = "匹配记录"
        ws.append(headers)

        # 简单设置列宽
        widths = [22, 36, 16, 10, 36, 20, 12, 12, 20, 20, 80]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = width

        wb.save(path)


def append_match_to_excel(
    path: str,
    client_id: str,
    my_name: str,
    my_sex: str,
    match_data: dict,
):
    """
    将 code=15 的匹配对象信息追加写入 Excel。
    """
    ensure_excel(path)

    wb = load_workbook(path)
    ws = wb["匹配记录"]

    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        client_id,
        my_name,
        my_sex,
        match_data.get("sel_userid", ""),
        match_data.get("sel_userNikename", ""),
        match_data.get("sel_userSex", ""),
        match_data.get("sel_userAge", ""),
        match_data.get("sel_userAddress", ""),
        match_data.get("sel_userTag", ""),
        json.dumps(match_data, ensure_ascii=False),
    ]

    ws.append(row)
    wb.save(path)


async def send_json(ws, payload: dict):
    """
    发送 JSON，并打印发送内容。
    """
    text = json.dumps(payload, ensure_ascii=False)
    print(f"\n[发送] {text}")
    await ws.send(text)


async def main():
    # 你可以按需修改这些身份字段
    client_id = gen_client_id()
    my_name = f"测试用户_{client_id[:6]}"
    my_sex = "男"

    sign_payload = {
        "act": "sign",
        "id": client_id,
        "name": my_name,
        "userSex": my_sex,
        "address_show": "false",
        "randomhealthmode": "0",
        "randomvipsex": "0",
        "randomvipaddress": "0",
        "userip": "",
        "useraddree": "测试地区",
        "randomvipcode": "",
    }

    random_payload = {
        "act": "random",
        "id": client_id,
        "preferSex": None,
        "preferAddress": False,
        "preferTag": None,
    }

    print("准备连接 WebSocket...")
    print("client_id:", client_id)
    print("nickname:", my_name)

    async with websockets.connect(WS_URL) as ws:
        print("\nWebSocket 已连接")

        # 1. sign
        await send_json(ws, sign_payload)

        signed = False
        matched = False
        matched_data = None

        start_time = datetime.now()

        while True:
            elapsed = (datetime.now() - start_time).total_seconds()
            if elapsed > MATCH_TIMEOUT:
                print(f"\n等待匹配超时：{MATCH_TIMEOUT} 秒")
                break

            try:
                raw = await asyncio.wait_for(ws.recv(), timeout=10)
            except asyncio.TimeoutError:
                print("\n[等待中] 暂无服务端消息...")
                continue

            print(f"\n[收到原始] {raw}")

            try:
                data = json.loads(raw)
            except json.JSONDecodeError:
                print("[解析失败] 非 JSON 消息")
                continue

            code = data.get("code")
            print("[解析后]", json.dumps(data, ensure_ascii=False, indent=2))

            # 2. sign 成功
            if code == 12:
                signed = True
                print("\n[状态] sign 成功，准备发送 random")
                await send_json(ws, random_payload)

            # 3. 进入匹配队列
            elif code == 13:
                print("\n[状态] 已进入匹配队列")

            # 4. 匹配成功
            elif code == 15:
                matched = True
                matched_data = data

                print("\n========== 匹配成功 ==========")
                print("对方ID:", data.get("sel_userid"))
                print("对方昵称:", data.get("sel_userNikename"))
                print("对方性别:", data.get("sel_userSex"))
                print("对方年龄:", data.get("sel_userAge"))
                print("对方地址:", data.get("sel_userAddress"))
                print("对方标签:", data.get("sel_userTag"))
                print("================================")

                append_match_to_excel(
                    path=EXCEL_PATH,
                    client_id=client_id,
                    my_name=my_name,
                    my_sex=my_sex,
                    match_data=data,
                )

                print(f"\n已写入 Excel：{EXCEL_PATH}")

                if CLOSE_AFTER_LOG:
                    random_out_payload = {
                        "act": "randomOut",
                        "id": client_id,
                        "msg": data.get("sel_userid"),
                    }
                    await send_json(ws, random_out_payload)
                    print("\n已发送 randomOut，释放当前匹配。")

                break

            # 5. 私聊消息推送/回显，仅打印，不回复
            elif code == 7:
                print("\n[消息] 收到私聊消息/消息回显")
                from_user = data.get("fromuser", {})
                print("发送方:", from_user.get("nickname"), from_user.get("id"))
                print("内容:", from_user.get("content"))

            # 6. 用户离线
            elif code == 3:
                print("\n[状态] 用户离线:", data.get("userIdMd5"))

            # 7. 状态类消息
            elif code in (18, 19):
                print("\n[状态类消息]", data)

            else:
                print("\n[其他 code]", code, data)

        print("\n程序结束")


if __name__ == "__main__":
    asyncio.run(main())